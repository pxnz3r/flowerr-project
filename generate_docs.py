import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_doc():
    doc = Document()

    # --- Styles ---
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    # --- Title Page ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run("Generative Study: Organic Flow")
    run.bold = True
    run.font.size = Pt(28)
    
    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle_p.add_run("Project Technical Documentation")
    run.font.size = Pt(18)
    
    doc.add_page_break()

    # --- Executive Summary ---
    doc.add_heading('1. Executive Summary', level=1)
    doc.add_paragraph(
        "Generative Study: Organic Flow is a web-based generative art installation implemented as a "
        "React-TypeScript application. The project explores the procedural growth of organic shapes "
        "(stems and flowers) using canvas-based animation and advanced SVG filters to emulate the "
        "texture of watercolor on paper."
    )

    # --- Project Architecture ---
    doc.add_heading('2. Project Architecture', level=1)
    doc.add_paragraph(
        "The application is built using modern web standards and follows a modular component-based "
        "architecture facilitated by React and Vite."
    )
    
    doc.add_heading('2.1 Technical Stack', level=2)
    doc.add_paragraph("• React 19 (Functional Components & Hooks)", style='List Bullet')
    doc.add_paragraph("• TypeScript (Type Safety & Scalability)", style='List Bullet')
    doc.add_paragraph("• Vite (Build Tool & Dev Server)", style='List Bullet')
    doc.add_paragraph("• Tailwind CSS v4 (Styling & Design System)", style='List Bullet')
    doc.add_paragraph("• Canvas API (Real-time Generative Rendering)", style='List Bullet')

    doc.add_heading('2.2 Component Breakdown', level=2)
    doc.add_paragraph(
        "The application is divided into four primary functional components:"
    )
    doc.add_paragraph("1. GenerativeArt: Handles the procedural logic, animation loop, and canvas rendering.", style='List Number')
    doc.add_paragraph("2. SVGFilters: Defines the watercolor and paper-noise filters used to post-process the canvas.", style='List Number')
    doc.add_paragraph("3. UILayer: Manages the typography and interactive textual elements (title and poem).", style='List Number')
    doc.add_paragraph("4. PaperTexture: A global overlay providing visual noise and texture.", style='List Number')

    # --- Implementation Details ---
    doc.add_heading('3. Implementation Details', level=1)
    
    doc.add_heading('3.1 Procedural Growth Algorithm', level=2)
    doc.add_paragraph(
        "The growth algorithm simulates botanical structures using randomized parametric paths. "
        "Stems are generated as a series of segments with localized 'noise' to create organic curves. "
        "Flowers are componentized into distinct types, each with its own petal distribution and "
        "blooming logic."
    )
    
    doc.add_heading('3.2 Visual Post-Processing', level=2)
    doc.add_paragraph(
        "To achieve a hand-painted look, the application utilizes CSS filters combined with SVG "
        "turbulence. The 'watercolor' filter uses fractal noise to displace the color of the canvas "
        "strokes, creating soft, bleeding edges characteristic of real watercolor pigments."
    )

    # --- Setup and Installation ---
    doc.add_heading('4. Setup and Installation', level=1)
    doc.add_paragraph("To run the project locally, ensure you have Node.js installed, then follow these steps:")
    doc.add_paragraph("1. Install dependencies: npm install", style='List Number')
    doc.add_paragraph("2. Start the development server: npm run dev", style='List Number')
    doc.add_paragraph("3. Access the application at http://localhost:5173", style='List Number')

    # --- Conclusion ---
    doc.add_heading('5. Conclusion', level=1)
    doc.add_paragraph(
        "This project successfully combines procedural generation with high-fidelity visual design, "
        "demonstrating the power of modern web technologies in creating immersive artistic experiences."
    )

    # --- Save ---
    doc_path = os.path.join(os.getcwd(), 'Project_Documentation.docx')
    doc.save(doc_path)
    print(f"Documentation saved to: {doc_path}")

if __name__ == "__main__":
    create_doc()
